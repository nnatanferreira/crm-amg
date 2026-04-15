import streamlit as st
from streamlit_gsheets import GSheetsConnection
import cloudinary
import cloudinary.uploader
import pandas as pd
import requests
import time
import re
from docx import Document
from io import BytesIO

# --- 1. CONFIGURAÇÕES CLOUDINARY ---
cloudinary.config( 
  cloud_name = "dybos073q", 
  api_key = "332571224149431", 
  api_secret = "bVdxjWZBAQ4_iGlaJAg0Kze3ZRU",
  secure = True
)

st.set_page_config(page_title="CRM AMG Multimarcas", page_icon="🚗", layout="wide")

# --- 2. FUNÇÕES DE APOIO ---
def formatar_para_br(valor):
    if not valor or str(valor).lower() in ["nan", "none", ""]: return "0,00"
    try:
        if "R$" in str(valor): return str(valor).replace("R$", "").strip()
        texto = str(valor).replace('R$', '').strip()
        apenas_numeros = re.sub(r'[^\d]', '', texto.split(',')[0].split('.')[0])
        v = float(apenas_numeros)
        return f"{v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
    except: return str(valor)

def limpar_id(valor):
    s = str(valor).strip()
    return s.split('.')[0] if '.' in s else s

def gerar_procuracao(r):
    try:
        doc = Document("modelo_procuracao.docx")
        
        substituicoes = {
            "{{NOME_TITULAR}}": str(r.get('nome_titular', '')),
            "{{CPF}}": str(r.get('tit_cpf', '')),
            "{{RG}}": str(r.get('tit_rg', '')),
            "{{RUA}}": str(r.get('tit_rua', '')),
            "{{NRUA}}": str(r.get('tit_num', '')),
            "{{BAIRRO}}": str(r.get('tit_bairro', '')),
            "{{COMPLEMENTO}}": str(r.get('tit_comp', '')),
            "{{CIDADE}}": str(r.get('tit_cid', '')),
            "{{UF}}": str(r.get('tit_est', '')),
            "{{PLACA}}": str(r.get('placa', '')),
            "{{CHASSI}}": str(r.get('chassi', '')),
            "{{RENAVAM}}": str(r.get('renavam', '')),
            "{{MARCAMODELO}}": f"{r.get('marca', '')}/{r.get('modelo', '')}",
            "{{ANO}}": str(r.get('ano', '')),
            "{{COR}}": str(r.get('cor', '')),
            "{{DATA}}": time.strftime("%d/%m/%Y")
        }

        for p in doc.paragraphs:
            for chave, valor in substituicoes.items():
                if chave in p.text:
                    p.text = p.text.replace(chave, valor)
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for chave, valor in substituicoes.items():
                        if chave in cell.text:
                            cell.text = cell.text.replace(chave, valor)

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer
    except:
        return None

# --- 3. CONEXÃO ---
conn = st.connection("gsheets", type=GSheetsConnection)

# --- 4. LOGIN ---
if "autenticado" not in st.session_state:
    st.markdown("<h1 style='text-align: center;'>🚗 CRM AMG</h1>", unsafe_allow_html=True)
    with st.form("login"):
        u = st.text_input("Usuário")
        p = st.text_input("Senha", type="password")
        if st.form_submit_button("ENTRAR"):
            if u == "amgmultimarcas" and p == "amg0031":
                st.session_state["autenticado"] = True
                st.rerun()
            else: st.error("Acesso negado")
else:
    menu = st.sidebar.radio("Navegação", ["➕ Cadastrar Veículo", "📑 Gerenciar Estoque"])
    if st.sidebar.button("Sair"):
        st.session_state.clear()
        st.rerun()

    # --- ABA: CADASTRAR ---
    if menu == "➕ Cadastrar Veículo":
        st.markdown("## 📝 Novo Cadastro")
        try:
            marcas = requests.get("https://fipe.parallelum.com.br/api/v2/cars/brands").json()
            dict_marcas = {m['name']: m['code'] for m in marcas}
            marca_n = st.selectbox("1. Marca", options=[""] + sorted(list(dict_marcas.keys())))
        except: marca_n = None

        if marca_n:
            modelos = requests.get(f"https://fipe.parallelum.com.br/api/v2/cars/brands/{dict_marcas[marca_n]}/models").json()
            dict_modelos = {m['name']: m['code'] for m in modelos}
            modelo_n = st.selectbox("2. Modelo", options=[""] + sorted(list(dict_modelos.keys())))

            if modelo_n:
                anos = requests.get(f"https://fipe.parallelum.com.br/api/v2/cars/brands/{dict_marcas[marca_n]}/models/{dict_modelos[modelo_n]}/years").json()
                dict_anos = {a['name']: a['code'] for a in anos}
                ano_sel = st.selectbox("3. Ano Modelo", options=[""] + list(dict_anos.keys()))

                if ano_sel:
                    dfipe = requests.get(f"https://fipe.parallelum.com.br/api/v2/cars/brands/{dict_marcas[marca_n]}/models/{dict_modelos[modelo_n]}/years/{dict_anos[ano_sel]}").json()
                    
                    with st.form("form_cadastro"):
                        st.subheader("🚗 Informações do Veículo")
                        c1, c2 = st.columns(2)
                        marca_v = c1.text_input("Marca", value=dfipe.get('brand'))
                        modelo_v = c1.text_input("Modelo", value=dfipe.get('model'))
                        placa_v = c1.text_input("Placa").upper()
                        foto_v = c1.file_uploader("📷 Foto Principal", type=['jpg','jpeg','png'])
                        
                        preco_v = c2.text_input("Preço de Venda", value=formatar_para_br(dfipe.get('price')))
                        km_v = c2.text_input("Quilometragem", value="0")
                        comb_v = c2.text_input("Combustível", value=dfipe.get('fuel'))
                        cor_v = c2.text_input("Cor")

                        st.markdown("---")
                        st.subheader("📑 Dados do CRLV")
                        tit_v = st.text_input("Nome Completo do Titular")
                        cc1, cc2 = st.columns(2)
                        ren_v = cc1.text_input("Renavam")
                        cha_v = cc2.text_input("Chassi").upper()
                        cc3, cc4 = st.columns(2)
                        tit_rg = cc3.text_input("RG do Titular")
                        tit_cpf = cc4.text_input("CPF do Titular")
                        
                        st.write("**Endereço do Titular:**")
                        rua_v = st.text_input("Rua/Logradouro")
                        bairro_v = st.text_input("Bairro")
                        cc5, cc6 = st.columns([1, 2])
                        num_v = cc5.text_input("Número")
                        comp_v = cc6.text_input("Complemento")
                        cc7, cc8 = st.columns(2)
                        cid_v = cc7.text_input("Cidade")
                        est_v = cc8.selectbox("UF", ["SP", "AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT", "MS", "MG", "PA", "PB", "PR", "PE", "PI", "RJ", "RN", "RS", "RO", "RR", "SC", "SE", "TO"])
                        
                        doc_v = st.file_uploader("📂 Documento (PDF/Imagem)", type=['pdf','jpg','jpeg','png'])

                        if st.form_submit_button("🚀 SALVAR"):
                            if not placa_v: st.error("Placa obrigatória!"); st.stop()
                            aviso = st.info("⏳ Salvando...")
                            url_img = cloudinary.uploader.upload(foto_v)['secure_url'] if foto_v else ""
                            url_doc = cloudinary.uploader.upload(doc_v)['secure_url'] if doc_v else ""
                            try: df_atual = conn.read(worksheet="Estoque", ttl=0).astype(str)
                            except: df_atual = pd.DataFrame()
                            
                            novo = pd.DataFrame([{
                                "marca": marca_v, "modelo": modelo_v, "placa": placa_v, "renavam": limpar_id(ren_v), "chassi": cha_v.strip(),
                                "cor": cor_v, "combustivel": comb_v, "preco": formatar_para_br(preco_v), "km": limpar_id(km_v), "foto": url_img,
                                "nome_titular": tit_v, "tit_rg": tit_rg, "tit_cpf": tit_cpf, "tit_rua": rua_v, "tit_bairro": bairro_v,
                                "tit_num": num_v, "tit_comp": comp_v, "tit_cid": cid_v, "tit_est": est_v, "doc_titular": url_doc, "ano": ano_sel
                            }])
                            conn.update(worksheet="Estoque", data=pd.concat([df_atual, novo], ignore_index=True).astype(str))
                            aviso.empty(); st.success("✅ Cadastrado!"); time.sleep(1); st.rerun()

    # --- ABA: ESTOQUE ---
    elif menu == "📑 Gerenciar Estoque":
        try:
            df = conn.read(worksheet="Estoque", ttl=0).dropna(how='all').astype(str)
        except: df = pd.DataFrame()

        if df.empty:
            st.info("Estoque vazio.")
        elif "edit_idx" in st.session_state:
            idx = st.session_state.edit_idx
            item = df.iloc[idx]
            with st.form("form_edicao"):
                st.subheader("✏️ Editar Veículo")
                c1, c2 = st.columns(2)
                m_e = c1.text_input("Marca", value=item['marca'])
                mo_e = c1.text_input("Modelo", value=item['modelo'])
                pl_e = c1.text_input("Placa", value=item['placa']).upper()
                f_v_e = c1.file_uploader("Trocar Foto")
                pr_e = c2.text_input("Preço", value=item['preco'])
                km_e = c2.text_input("KM", value=item['km'])
                comb_e = c2.text_input("Combustível", value=item.get('combustivel',''))
                cor_e = c2.text_input("Cor", value=item.get('cor',''))
                st.markdown("---")
                tit_e = st.text_input("Titular", value=item.get('nome_titular',''))
                ren_e = st.text_input("Renavam", value=item.get('renavam',''))
                cha_e = st.text_input("Chassi", value=item.get('chassi',''))
                rg_e = st.text_input("RG", value=item.get('tit_rg',''))
                cpf_e = st.text_input("CPF", value=item.get('tit_cpf',''))
                rua_e = st.text_input("Rua", value=item.get('tit_rua',''))
                bai_e = st.text_input("Bairro", value=item.get('tit_bairro',''))
                num_e = st.text_input("Nº", value=item.get('tit_num',''))
                com_e = st.text_input("Compl.", value=item.get('tit_comp',''))
                cid_e = st.text_input("Cidade", value=item.get('tit_cid',''))
                est_e = st.text_input("UF", value=item.get('tit_est',''))

                if st.form_submit_button("💾 SALVAR ALTERAÇÕES"):
                    df.at[idx, 'marca'] = m_e; df.at[idx, 'modelo'] = mo_e; df.at[idx, 'placa'] = pl_e
                    df.at[idx, 'preco'] = pr_e; df.at[idx, 'km'] = km_e; df.at[idx, 'nome_titular'] = tit_e
                    df.at[idx, 'renavam'] = ren_e; df.at[idx, 'chassi'] = cha_e; df.at[idx, 'tit_rg'] = rg_e
                    df.at[idx, 'tit_cpf'] = cpf_e; df.at[idx, 'tit_rua'] = rua_e; df.at[idx, 'tit_bairro'] = bai_e
                    df.at[idx, 'tit_num'] = num_e; df.at[idx, 'tit_comp'] = com_e; df.at[idx, 'tit_cid'] = cid_e; df.at[idx, 'tit_est'] = est_e
                    conn.update(worksheet="Estoque", data=df.astype(str))
                    st.success("✅ Atualizado!"); del st.session_state.edit_idx; time.sleep(1); st.rerun()
                if st.form_submit_button("❌ CANCELAR"):
                    del st.session_state.edit_idx; st.rerun()
        else:
            for i, r in df.iterrows():
                with st.container():
                    col1, col2 = st.columns([1, 2])
                    if r.get('foto'): col1.image(r['foto'], use_container_width=True)
                    col2.subheader(f"{r['marca']} {r['modelo']} - {r['placa']}")
                    col2.write(f"Titular: {r['nome_titular']}")
                    btn_c1, btn_c2, btn_c3 = col2.columns(3)
                    if btn_c1.button(f"✏️ Editar", key=f"e{i}"):
                        st.session_state.edit_idx = i; st.rerun()
                    
                    doc_ready = gerar_procuracao(r)
                    if doc_ready:
                        btn_c2.download_button("📜 Procuração", data=doc_ready, file_name=f"Procuracao_{r['placa']}.docx", key=f"p{i}")
                    else:
                        btn_c2.warning("Modelo .docx não encontrado")
                        
                    if btn_c3.button(f"🗑️ Excluir", key=f"d{i}"):
                        conn.update(worksheet="Estoque", data=df.drop(i).astype(str))
                        st.rerun()
                st.markdown("---")
